#!/usr/bin/env python3
"""从 Markdown 生成 Word 技术报告。"""

from pathlib import Path
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn

MD_PATH = Path("/Users/mingrun/resume-agent-system/docs/tuning_report_content.md")
OUT_PATH = Path("/Users/mingrun/resume-agent-system/docs/LangGraph多智能体任务调优报告.docx")


def set_run_font(run, name="Microsoft YaHei", size=10.5, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = color


def add_heading(doc, text, level=1):
    heading = doc.add_heading(level=level)
    size = 18 if level == 1 else (14 if level == 2 else 12)
    color = RGBColor(0, 0, 128) if level == 1 else None
    set_run_font(heading.add_run(text), size=size, bold=True, color=color)
    return heading


def add_paragraph(doc, text, bold=False, italic=False, size=10.5):
    p = doc.add_paragraph()
    set_run_font(p.add_run(text), size=size, bold=bold)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.ONE_POINT_FIVE
    p.paragraph_format.space_after = Pt(6)
    return p


def add_code_block(doc, lines):
    p = doc.add_paragraph()
    text = "\n".join(lines)
    run = p.add_run(text)
    run.font.name = "Consolas"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Consolas")
    run.font.size = Pt(9)
    p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
    p.paragraph_format.left_indent = Inches(0.2)
    p.paragraph_format.space_after = Pt(6)


def parse_table(lines):
    headers = [c.strip() for c in lines[0].split("|") if c.strip()]
    rows = []
    for line in lines[2:]:
        cells = [c.strip() for c in line.split("|")]
        cells = [c for c in cells if c]
        if cells:
            rows.append(cells)
    return headers, rows


def add_table(doc, headers, rows):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    hdr = table.rows[0].cells
    for i, h in enumerate(headers):
        hdr[i].text = h
        for run in hdr[i].paragraphs[0].runs:
            set_run_font(run, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = val
            for run in cells[i].paragraphs[0].runs:
                set_run_font(run)


def main():
    doc = Document()

    # 封面
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("LangGraph 多智能体文章写作任务\n调优过程技术报告"), size=22, bold=True, color=RGBColor(0, 0, 128))

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    set_run_font(p.add_run("模型：Kimi k2p5-coding | 任务：D-B-E-C 四阶段智能体平台\n报告日期：2026-07-13"), size=12)

    doc.add_page_break()

    lines = MD_PATH.read_text(encoding="utf-8").splitlines()
    i = 0
    in_code = False
    code_lines = []
    table_lines = []
    in_table = False

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, code_lines)
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        if line.startswith("|"):
            in_table = True
            table_lines.append(line)
            i += 1
            continue
        elif in_table:
            headers, rows = parse_table(table_lines)
            add_table(doc, headers, rows)
            table_lines = []
            in_table = False

        if line.startswith("# "):
            add_heading(doc, line[2:].strip(), level=1)
        elif line.startswith("## "):
            add_heading(doc, line[3:].strip(), level=2)
        elif line.startswith("### "):
            add_heading(doc, line[4:].strip(), level=3)
        elif line.strip().startswith("- ") or line.strip().startswith("1. ") or line.strip().startswith("a) "):
            add_paragraph(doc, line.strip(), size=10.5)
        elif line.strip():
            add_paragraph(doc, line.strip(), size=10.5)

        i += 1

    if in_table:
        headers, rows = parse_table(table_lines)
        add_table(doc, headers, rows)

    doc.save(OUT_PATH)
    print(f"报告已保存到: {OUT_PATH}")


if __name__ == "__main__":
    main()
